# SPDX-License-Identifier: Apache-2.0
"""Regression tests for the generated bench l2 geometry delivery guide."""

# Standard
import importlib.util
from pathlib import Path

# Third Party
from docx import Document


REPO_ROOT = Path(__file__).resolve().parents[2]
GENERATOR = REPO_ROOT / "scripts/ipu-poc/build_geometry_delivery_doc.py"


def generated_text(output: Path) -> str:
    """Build the guide and return every generated paragraph and table cell."""
    spec = importlib.util.spec_from_file_location("geometry_delivery", GENERATOR)
    assert spec is not None
    assert spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    module.build_document(output)

    doc = Document(output)
    paragraphs = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(paragraphs)


def test_guide_walks_a_new_user_through_current_sweep(tmp_path: Path) -> None:
    """The guide exposes the current inventory, sweep, results, and caveats."""
    text = generated_text(tmp_path / "guide.docx")

    assert "Use the supplied inventory file" in text
    assert "scripts/ipu-poc/inventories/mmg-two-initiator.env" in text
    assert "five page-burst profiles" in text
    assert "Mixtral 128 KiB" in text
    assert "Mixtral 512 KiB" in text
    assert "BENCH_MOUNT/bench-l2/<run-id>/<hostname>" in text
    assert "/mnt/lmcache-stage2/bench-l2" not in text
    assert "Before you run preflight" in text
    assert "Collect the result JSON" in text
    assert 'ssh "$host" bash -s -- "$RUN_ID"' in text
    assert 'for repo in "$HOME/LMCache" /root/LMCache; do' in text
    assert 'find "$repo/results" -maxdepth 1 -name "$run_id-*.json" -print' in text
    assert "IN_FLIGHTS='8 16 24'" in text
    assert "WARMUP_SEC=60 DURATION_SEC=120 INCLUDE_MINIMAX=1" in text
    assert "O_DIRECT-compatible padding profile" in text
    assert "buffered I/O" not in text
    assert "64 KiB FIO cells control only the 64 KiB profile" in text
    assert "Optional telemetry setup" in text
