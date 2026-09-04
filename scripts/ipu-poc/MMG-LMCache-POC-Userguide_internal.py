# SPDX-License-Identifier: Apache-2.0
"""Build the "KV Cache Offload over Falcon" operator guide.

Needs python-docx, which is not a project dependency: `pip install python-docx`.
"""

# Standard
import argparse
from pathlib import Path

# Third Party
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REPO_ROOT = Path(__file__).resolve().parents[2]
DIAGRAMS_DIR = REPO_ROOT / "docs/design/v1/platform/ipu-poc/diagrams/brief"


def shade_cell(cell, fill: str) -> None:
    """Apply a background fill to a table cell."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def pad_cell(cell, twips: int) -> None:
    """Apply equal internal padding to a table cell."""
    margins = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        margin = OxmlElement(f"w:{side}")
        margin.set(qn("w:w"), str(twips))
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    cell._tc.get_or_add_tcPr().append(margins)


def prevent_row_split(row) -> None:
    """Keep a one-row command block on one page."""
    properties = row._tr.get_or_add_trPr()
    no_split = OxmlElement("w:cantSplit")
    properties.append(no_split)


def add_bottom_rule(paragraph) -> None:
    """Add the title-block divider beneath a paragraph."""
    paragraph_properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), "BBBBBB")
    borders.append(bottom)
    paragraph_properties.append(borders)


def repeat_table_header(row) -> None:
    """Repeat a table header when the table crosses a page boundary."""
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def command(doc: Document, text: str) -> None:
    """Add a compact shaded command block."""
    table = doc.add_table(rows=1, cols=1)
    prevent_row_split(table.rows[0])
    cell = table.cell(0, 0)
    shade_cell(cell, "F1F2F4")
    pad_cell(cell, 120)

    lines = text.strip().splitlines()
    paragraph = cell.paragraphs[0]
    paragraph.style = "Command"
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run("\n".join(lines))


REPO_BLOB_BASE = "https://github.com/tsg-/LMCache/blob/feat/bench-l2-geometry-handoff/"

COLLECTOR_FILES = [
    (
        "docs/design/v1/platform/ipu-poc/instrumentation/host/bin/"
        "rdma_hwcounters_textfile.sh",
        "irdma hw_counters -- directional RDMA operation counters",
    ),
    (
        "docs/design/v1/platform/ipu-poc/instrumentation/host/bin/"
        "rdma_nic_textfile.sh",
        "ethtool -S fabric-NIC counters -- traffic-presence diagnostic",
    ),
    (
        "docs/design/v1/platform/ipu-poc/instrumentation/host/bin/"
        "acc_telemetry_textfile.py",
        "ACC gRPC RC byte counters -- authoritative Falcon payload-byte source",
    ),
    (
        "docs/design/v1/platform/ipu-poc/instrumentation/host/bin/"
        "nvme_stats_textfile.sh",
        "NVMe SMART per namespace",
    ),
    (
        "docs/design/v1/platform/ipu-poc/instrumentation/host/bin/"
        "pcm_memory_textfile.sh",
        "Intel PCM DRAM read/write bandwidth per socket",
    ),
    (
        "docs/design/v1/platform/ipu-poc/instrumentation/host/bin/"
        "numa_stats_textfile.sh",
        "Kernel node memory, allocation, and CPU-time counters",
    ),
    (
        "docs/design/v1/platform/ipu-poc/instrumentation/host/bin/"
        "pcm_pcie_textfile.sh",
        "Intel PCM PCIe/DDIO bandwidth per socket (mmgt only)",
    ),
    (
        "docs/design/v1/platform/ipu-poc/instrumentation/host/bin/"
        "mmgt_nic_textfile.sh",
        "ethtool -S on both fabric ports (mmgt only, replaces rdma_nic)",
    ),
    (
        "scripts/ipu-poc/acc_ssh_stats.py",
        "ACC core usage + tele_cli -t global, over SSH via the IMC",
    ),
    (
        "scripts/ipu-poc/acc_grpc_tunnel.py",
        "Persistent target-only IMC -> ACC gRPC tunnel and protobuf staging",
    ),
    (
        "scripts/ipu-poc/install_acc_stats.sh",
        "Idempotent installer for the two collectors above, per host",
    ),
]


def add_hyperlink(paragraph, url: str, text: str, size_pt: float = 9.5) -> None:
    """Insert a real, clickable hyperlink run into *paragraph*.

    python-docx has no built-in hyperlink API; this is the standard
    low-level workaround -- register the URL as an external relationship
    on the document part, then build the <w:hyperlink> XML by hand. The
    run this creates is nested inside <w:hyperlink>, not a direct child
    of <w:p>, so python-docx's own ``paragraph.runs`` will not see it --
    set formatting here, in the XML, rather than through the Run API
    afterward.
    """
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(int(size_pt * 2)))
    run_properties.append(color)
    run_properties.append(underline)
    run_properties.append(size)
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def collector_table(doc: Document) -> None:
    """List every new collector, linked to its source in the tsg fork."""
    table = doc.add_table(rows=1 + len(COLLECTOR_FILES), cols=2)
    table.style = "Table Grid"
    for row in table.rows:
        prevent_row_split(row)
    repeat_table_header(table.rows[0])

    for col, text in enumerate(("Collector source", "Purpose")):
        cell = table.cell(0, col)
        shade_cell(cell, "D9EAF7")
        pad_cell(cell, 40)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        run.font.size = Pt(9)

    for r, (path, purpose) in enumerate(COLLECTOR_FILES, start=1):
        path_cell = table.cell(r, 0)
        pad_cell(path_cell, 40)
        path_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        path_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if r % 2 == 0:
            shade_cell(path_cell, "F1F2F4")
        add_hyperlink(path_cell.paragraphs[0], REPO_BLOB_BASE + path, path, size_pt=8.5)

        purpose_cell = table.cell(r, 1)
        pad_cell(purpose_cell, 40)
        purpose_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        purpose_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if r % 2 == 0:
            shade_cell(purpose_cell, "F1F2F4")
        run = purpose_cell.paragraphs[0].add_run(purpose)
        run.font.size = Pt(8.5)
    doc.add_paragraph()


def figure(
    doc: Document, image_path: Path, caption: str, width_in: float = 6.3
) -> None:
    """Center a captioned diagram, scaled to fit the text column."""
    picture_paragraph = doc.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.add_run().add_picture(str(image_path), width=Inches(width_in))
    caption_paragraph = doc.add_paragraph(caption, style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def data_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    """Add a shaded-header data table sized to its content."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in table.rows:
        prevent_row_split(row)

    for col, text in enumerate(headers):
        cell = table.cell(0, col)
        shade_cell(cell, "D9EAF7")
        pad_cell(cell, 40)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        run.font.size = Pt(9)

    for r, row in enumerate(rows, start=1):
        for col, text in enumerate(row):
            cell = table.cell(r, col)
            pad_cell(cell, 40)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if r % 2 == 0:
                shade_cell(cell, "F1F2F4")
            run = cell.paragraphs[0].add_run(text)
            run.font.size = Pt(9)
    doc.add_paragraph()


def build_styles(doc: Document) -> None:
    """Set the compact operator-guide type scale and rhythm."""
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, before, after in (
        ("Title", 0, 3),
        ("Subtitle", 0, 14),
        ("Heading 1", 15, 4),
        ("Heading 2", 10, 3),
    ):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for name in ("Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor(0x30, 0x6F, 0xBF)

    caption = doc.styles["Caption"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(10)

    command_style = doc.styles.add_style("Command", 1)
    command_style.base_style = doc.styles["Normal"]
    command_style.font.name = "Menlo"
    command_style.font.size = Pt(8.5)
    command_style.paragraph_format.space_after = Pt(0)
    command_style.paragraph_format.line_spacing = 1.2


def build_document(output: Path) -> None:
    """Create the DOCX guide at ``output``."""
    doc = Document()
    build_styles(doc)
    for section in doc.sections:
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(3)
    title_run = title.add_run("KV Cache Offload over Falcon")
    title_run.font.name = "Aptos Display"
    title_run.font.size = Pt(26)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor(0x30, 0x6F, 0xBF)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle_run = subtitle.add_run(
        "Running LMCache bench l2 in a larger MMG-400 setup"
    )
    subtitle_run.font.name = "Aptos"
    subtitle_run.font.size = Pt(14)
    subtitle_run.italic = True
    subtitle_run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    add_bottom_rule(subtitle)
    doc.add_paragraph(
        "This guide demonstrates the MMG-400 two-initiator setup (mmgi0 and "
        "mmgi1) against one target. The inventory names machines only. The "
        "coordinator rejects an unmounted or local filesystem instead of "
        "accepting a user-supplied storage path. The tooling accepts more "
        "initiators, but this guide's topology, procedure, and results are "
        "for two."
    )
    doc.add_heading("PoC Scope", level=1)
    for scope_item in (
        "Intel IPU Falcon transport moves KV-cache traffic at target line "
        "rate, and the cache retrieve and store paths behave as "
        "representative NVMe-oF flows.",
        "The Xeon-based storage setup has enough compute, memory, and I/O "
        "bandwidth, and can use features such as DDIO to deliver line-rate "
        "throughput.",
    ):
        doc.add_paragraph(scope_item, style="List Bullet")
    doc.add_heading("Test Setup", level=1)
    doc.add_paragraph(
        "Traffic generators (1x 400 GbE each, 2x 400 GbE combined) and one "
        "KV Cache Target, connected over a Falcon fabric. "
        "The target serves eight 1.92 TB Solidigm D7-PS1010 Gen5 x4 NVMe "
        "SSDs over 4x 400 GbE, as built. The full-scale design point is 16x "
        "Gen5 x4 NVMe SSDs on the same 4x 400 GbE target."
    )
    figure(
        doc,
        DIAGRAMS_DIR / "brief-mmg400-setup-asbuilt.png",
        "The rig as built: two initiators/traffic generators, one target, "
        "8 drives.",
    )

    doc.add_heading("System configuration", level=2)
    doc.add_paragraph(
        "For consistent performance, set BIOS System Profile to Performance, "
        "disable C-States, and make sure the OS power/performance governor is "
        'set to "performance".'
    )

    doc.add_heading("Benchmark Definition", level=1)
    doc.add_heading("KV Caching, LMCache and the Storage (L2) Tier", level=2)
    doc.add_paragraph(
        "A model creates key/value (KV) tensors while processing a prompt; "
        "reusing them avoids recomputing the prompt's attention work. LMCache "
        "is the open-source KV cache engine under test — it sits between the "
        "inference engine (e.g., vLLM) and GPU memory, offloading KV tensors "
        "that don't fit in GPU HBM so they can be reused instead of "
        "recomputed. It organizes that offloaded data into tiers: a fast local "
        "tier holds recently-active objects, while a backing tier holds "
        "objects that must be fetched before reuse. In this guide, L2 means "
        "that backing tier. The bench l2 command exercises LMCache's L2 "
        "adapter directly; on this rig the adapter is fs_native on the XFS "
        "filesystem mounted over NVMe-oF. It measures cache-object I/O, not "
        "end-to-end model inference."
    )
    doc.add_heading("Why LMCache bench l2, not FIO alone", level=2)
    doc.add_paragraph(
        "This sweep first writes a small model-shaped corpus, then runs a "
        "timed, read-only sustained load over that corpus. The store phase "
        "prepares the keys; it is not concurrent with, or included in, the "
        "timed read result. FIO is a controlled storage-and-fabric comparator "
        "(Appendix E), while bench l2 adds model geometry, object grouping, "
        "and adapter behavior. Neither result is an end-to-end inference "
        "throughput measurement."
    )
    doc.add_paragraph(
        "bench l2 also has a separate mixed mode. The 5:1 read/write "
        "characterization below illustrates that mode's millisecond-scale "
        "overlap and burst behavior; it is not the timed workload in Step 4."
    )
    timeline_lead = doc.add_paragraph()
    timeline_lead.add_run(
        "Measured LMCache bench l2 timeline — reader/writer sequencing"
    ).bold = True
    doc.add_paragraph(
        "Path: mkp1 → fs_native → XFS on md0 RAID0 → Falcon-backed NVMe-oF "
        "→ mkp2 physical NVMe SSDs. Result from E2E Falcon/NVMe-oF "
        "measurement."
    )
    doc.add_paragraph(
        "Geometry: 256 KiB KV pages; 56 objects per submit; 14 MiB per "
        "submit; O_DIRECT; 16 workers; 8 submits in flight. The 28 GiB read "
        "corpus was prepopulated. Stores used a fresh prefix for the run."
    )
    figure(
        doc,
        DIAGRAMS_DIR / "brief-bench-l2-100ms-window.png",
        "5:1 mixed mode: B = a read and write both fired that millisecond; "
        "R = read only.",
    )
    doc.add_paragraph("Write bursts recur with a regular distribution, not uniformly:")
    figure(
        doc,
        DIAGRAMS_DIR / "brief-bench-l2-write-burst-histogram.png",
        "5:1 mixed mode: how long each write burst runs before going quiet.",
    )
    doc.add_heading("Model coverage", level=2)
    data_table(
        doc,
        ["Model", "Attention", "Object / page size", "Geometry type", "Notes"],
        [
            ["Mixtral 8x22B FP8", "GQA", "64 KiB", "Page-burst", ""],
            ["Mixtral 8x22B FP8", "GQA", "128 KiB", "Page-burst", ""],
            ["DeepSeek-V3 FP8", "MLA", "144 KiB", "Page-burst", ""],
            ["Mixtral 8x22B FP8", "GQA", "256 KiB", "Page-burst", ""],
            ["Mixtral 8x22B FP8", "GQA", "512 KiB", "Page-burst", ""],
            [
                "MiniMax-M3 bf16 KV / fp8 indexer, TP=8",
                "DSA",
                "~9.34 MiB",
                "Object-group",
                "O_DIRECT padded; indexer assumption unconfirmed",
            ],
        ],
    )
    doc.add_paragraph(
        "Mixtral variants isolate page size as the single variable. MiniMax "
        "packs 60-layer K/V and 57-layer DSA indexer into one object per "
        "(chunk, kv rank), a shape no uniform page size can express."
    )
    doc.add_heading("From model profile to bench l2 parameters", level=2)
    doc.add_paragraph(
        "Worked example: the Mixtral 256 KiB row above, from architecture "
        "YAML to object size to the wrapper invocation the sweep runs "
        "underneath."
    )
    command(
        doc,
        """architecture:
  num_layers: 56
  num_kv_heads: 8
  head_size: 128
  kv_size: 2
quantization:
  dtype_bytes: 1  # FP8
chunking:
  tokens_per_chunk: 128""",
    )
    doc.add_paragraph(
        "2 KV x 8 heads x 128 values x 1 byte x 128 tokens = 262,144 bytes "
        "= 256 KiB per object. 56 layer objects x 256 KiB = 14 MiB per "
        "submit."
    )
    command(
        doc,
        """PROFILE=scripts/ipu-poc/models/mixtral_8x22b_fp8.yaml
BASE_PATH=/mnt/lmcache/bench-l2/example/$(hostname -s) \\
    PREFIX=example-$(hostname -s)-if16 \\
    IN_FLIGHT=16 \\
    ROUNDS=2 \\
    bash scripts/ipu-poc/run_model_geometry.sh store "$PROFILE"
BASE_PATH=/mnt/lmcache/bench-l2/example/$(hostname -s) \\
    PREFIX=example-$(hostname -s)-if16 \\
    IN_FLIGHT=16 \\
    WARMUP_SEC=60 \\
    DURATION_SEC=120 \\
    METRICS_PORT=9101 \\
    OUTPUT=results/example-$(hostname -s)-mixtral_8x22b_fp8-if16.json \\
    bash scripts/ipu-poc/run_model_geometry.sh sustained-load "$PROFILE" """,
    )
    doc.add_paragraph(
        "This copyable one-model example uses the wrapper, which adds the "
        "profile hash to the key namespace. Step 4 repeats the same store "
        "then sustained-load sequence for every (profile, in-flight) cell."
    )

    doc.add_heading("Execution assumptions", level=1)
    doc.add_paragraph(
        "Before beginning: the coordinator needs passwordless SSH to both "
        "initiators and the target; every initiator needs an LMCache checkout "
        "at $HOME/LMCache or /root/LMCache; and the NVMe-oF filesystem must "
        "be mounted at /mnt/lmcache unless BENCH_MOUNT is explicitly set. "
        "Run every coordinator command with "
        "scripts/ipu-poc/inventories/mmg-two-initiator.env. No dedicated "
        "controller host is required; preflight checks these conditions "
        "before the sweep writes data."
    )

    doc.add_heading("0. Clone the custom LMCache fork", level=1)
    command(
        doc,
        """git clone https://github.com/tsg-/LMCache.git ~/LMCache
cd ~/LMCache
git checkout feat/bench-l2-geometry-handoff
git rev-parse HEAD""",
    )
    doc.add_paragraph(
        "This fork of upstream LMCache adds bench l2 geometry profiles for "
        "Mixtral 8x22B FP8 (64, 128, 256, and 512 KiB pages), DeepSeek-V3 "
        "FP8 (144 KiB MLA), and MiniMax-M3 (bf16 KV / fp8 indexer with an "
        "O_DIRECT-padded variant). It also includes the geometry harness, "
        "corpus-identity checks, sweep tooling, and instrumentation fixes "
        "used by this guide. Clone to this path on the coordinator and on "
        "every initiator: the remote wrappers look for the checkout at "
        "$HOME/LMCache (or /root/LMCache) and abort if they don't find it "
        "there. Run subsequent commands from LMCache/. Record the printed "
        "commit with the result JSON; the branch can advance after this "
        "guide is generated."
    )

    doc.add_heading("1. Name your hosts in the inventory file", level=1)
    doc.add_paragraph(
        "Env file scripts/ipu-poc/inventories/mmg-two-initiator.env defines "
        "host names for the test rig."
    )
    command(
        doc,
        """# scripts/ipu-poc/inventories/mmg-two-initiator.env
INITIATOR_HOSTS=(mmgi0 mmgi1)
TARGET_HOST=mmgt""",
    )
    doc.add_paragraph(
        "The coordinator runs the benchmark on each initiator. The target "
        "remains storage-only, contacted only to verify NVMe-oF export "
        "addresses."
    )
    doc.add_heading("2. Validate the environment (preflight)", level=1)
    command(
        doc,
        """bash scripts/ipu-poc/run_geometry_inventory.sh preflight \\
    scripts/ipu-poc/inventories/mmg-two-initiator.env""",
    )
    doc.add_paragraph(
        "Preflight checks the checkout, interpreter, mount, and NVMe-oF "
        "controller on every initiator, and stops before writing anything. "
        "It confirms the environment can run a sweep; it does not establish "
        "this rig's bandwidth ceiling. Do that once with Appendix E before "
        "trusting a result."
    )

    doc.add_heading("3. Confirm store/read correctness (identity gate)", level=1)
    command(
        doc,
        """RUN_ID=mmg-verify-$(date +%Y%m%d-%H%M%S) \\
    bash scripts/ipu-poc/run_geometry_inventory.sh verify \\
    scripts/ipu-poc/inventories/mmg-two-initiator.env""",
    )
    doc.add_paragraph(
        "The verification gate writes a small corpus, requires the storing "
        "profile to read it back, and requires a same-geometry mismatch to "
        "read zero objects. It is a correctness check, not a bandwidth "
        "result"
    )

    doc.add_heading("4. Run the benchmark sweep", level=1)
    command(
        doc,
        """export RUN_ID=mmg-geometry-$(date +%Y%m%d-%H%M%S)
mkdir -p logs
nohup env \\
    IN_FLIGHTS='8 16 24' \\
    WARMUP_SEC=60 \\
    DURATION_SEC=120 \\
    INCLUDE_MINIMAX=1 \\
    bash scripts/ipu-poc/run_geometry_inventory.sh sweep \\
    scripts/ipu-poc/inventories/mmg-two-initiator.env \\
  >"logs/$RUN_ID-sweep.log" 2>&1 &
echo "coordinator pid=$!  log=logs/$RUN_ID-sweep.log"
tail -f "logs/$RUN_ID-sweep.log" """,
    )
    doc.add_paragraph(
        "Each initiator runs the same sequence at 8, 16, and 24 in-flight "
        "submits. At each setting, it stores and then reads the page-burst "
        "profiles in page-size order, followed by MiniMax-M3. The two "
        "initiators run in parallel, but each initiator runs its cells "
        "serially."
    )
    doc.add_paragraph(
        "A model's three in-flight measurements are separated by a pass over "
        "the other profiles; they are not back-to-back. Treat MiniMax-M3 "
        "cautiously until its fp8 DSA-indexer geometry assumption is "
        "confirmed."
    )
    doc.add_paragraph(
        "Each host writes under BENCH_MOUNT/bench-l2/{run-id}/{hostname} "
        "(default /mnt/lmcache) with a host-specific key prefix. Result JSON "
        "files are written to the remote checkout's results/ directory. "
        "Timed windows alone take at least 54 minutes per host, plus stores "
        "and setup."
    )

    doc.add_heading("5. Review results", level=1)
    doc.add_paragraph(
        "For each result JSON, require total_success == total_keys. A miss "
        "invalidates the bandwidth figure. The default corpus is small enough "
        "to fit in DRAM on most benchmark hosts, so it is a functional sweep; "
        "scale the corpus beyond memory before making a storage-performance "
        "claim."
    )
    doc.add_heading("Optional: archive result JSON on the coordinator", level=2)
    doc.add_paragraph(
        "Each initiator writes 18 JSON files to its LMCache results/ "
        "directory, 36 total. Keep the exported RUN_ID from step 4. Use this "
        "optional command only when a coordinator-side archive is needed; it "
        "copies the files and aborts if either initiator did not produce 18."
    )
    command(
        doc,
        """set -euo pipefail
source scripts/ipu-poc/inventories/mmg-two-initiator.env
mkdir -p "collected-results/$RUN_ID"
for host in "${INITIATOR_HOSTS[@]}"; do
  mkdir -p "collected-results/$RUN_ID/$host"
  ssh "$host" bash -s -- "$RUN_ID" <<'REMOTE' | \\
    tar -xzf - -C "collected-results/$RUN_ID/$host"
set -euo pipefail
run_id=$1
for repo in "$HOME/LMCache" /root/LMCache; do
  [ -x "$repo/scripts/ipu-poc/run_model_geometry.sh" ] && break
done
[ -x "$repo/scripts/ipu-poc/run_model_geometry.sh" ] ||
  { echo "ABORT: LMCache geometry checkout not found" >&2; exit 2; }
mapfile -t files < <(
  find "$repo/results" -maxdepth 1 -name "$run_id-*.json" -exec basename {} ';' | sort
)
[ "${#files[@]}" -eq 18 ] ||
  { echo "ABORT: expected 18 result JSON files, found ${#files[@]}" >&2; exit 2; }
tar -C "$repo/results" -czf - -- "${files[@]}"
REMOTE
done""",
    )

    doc.add_page_break()
    doc.add_heading("Reference results and appendices", level=1)
    doc.add_heading("Established bandwidth limits — 2x 400 GbE IPU setup", level=2)
    doc.add_paragraph(
        "FIO establishes the envelope this path can move bytes at, "
        "independent of any application (see Appendix E). The tables below "
        "cover 64 KiB and 256 KiB locally, and 64 KiB through 1 MiB over the "
        "remote fs_native path."
    )
    doc.add_heading("Peak local FIO throughput (on mmgt, before the fabric)", level=2)
    data_table(
        doc,
        ["Block size", "Peak (Gb/s)", "Peak (GB/s)", "Outstanding at peak"],
        [
            ["64 KiB", "928", "116", "512 (8 jobs × qd64)"],
            ["256 KiB", "952", "119", "256 (8 jobs × qd32)"],
        ],
    )
    doc.add_paragraph(
        "Measured 2026-09-01 after the Falcon rebuild, direct=1, libaio, 8 "
        "jobs (one per Solidigm namespace), randread against the raw "
        "/dev/nvmeXn1 devices, before nvmet export and before the fabric. "
        "~14.5-14.9 GB/s per drive, in line with the D7-PS1010 brief's "
        "up-to-14.5 GB/s sequential read spec. This is the media and PCIe "
        "ceiling the remote figures below are measured against."
    )
    doc.add_heading(
        "Peak remote FIO throughput (md0 + XFS over NVMe-oF — the fs_native "
        "envelope)",
        level=2,
    )
    data_table(
        doc,
        [
            "Block size",
            "Combined (Gb/s)",
            "Per initiator (Gb/s)",
            "Outstanding at peak",
        ],
        [
            ["64 KiB", "651.2", "~325.6", "256 (8 jobs × qd32)"],
            ["256 KiB", "768.0", "~384.0", "512 (8 jobs × qd64)"],
            ["512 KiB", "768.0", "~384.0", "128 (8 jobs × qd16)"],
            ["1 MiB", "768.8", "~384.4", "128 (8 jobs × qd16)"],
        ],
    )
    doc.add_paragraph(
        "Re-measured 2026-09-01 after mmgt/mmgi0/mmgi1 were rebooted onto a "
        "new Falcon setup and the storage stack (nvmet export, NVMe-oF "
        "connect, md0 RAID0, XFS) was rebuilt from scratch on both "
        "initiators -- both driven in lockstep, direct=1, libaio, 8 jobs, "
        "128 GB real-data corpus per host (0 unwritten extents, verified), "
        "both links at 400000 Mb/s / MTU 9100. This supersedes the "
        "2026-08-29 pre-reboot figures (681 / 677 Gb/s): 256 KiB now measures "
        "notably higher (768.0 vs 677 Gb/s) and 64 KiB somewhat lower (651.2 "
        "vs 681 Gb/s) on the rebuilt fabric -- reported as measured rather "
        "than reconciled to the earlier run. The listed remote peak occurs "
        "at 256 aggregate outstanding for 64 KiB and 512 for 256 KiB. At "
        "512 KiB and 1 MiB it occurs at 128 outstanding, consistent with "
        "larger I/O sizes reaching the byte-rate limit at a shallower queue "
        "depth. Compare a bench l2 result against the FIO cell nearest its "
        "actual concurrency and block size, not a single peak number."
    )
    doc.add_paragraph(
        "256 KiB through 1 MiB I/O sizes yield the same maximum peak (384 "
        "Gb/s per initiator, 96% of 400GbE wire rate). This is expected: "
        "RDMA path MTU is capped at 4096 B by specification, so larger block "
        "sizes produce more packets per I/O without improving per-packet "
        "efficiency."
    )

    doc.add_heading("5:1 mixed FIO results (placeholder)", level=2)
    doc.add_paragraph(
        "No 5:1 mixed FIO result is published yet. Populate this table only "
        "with a separately run 5:1 read/write FIO workload."
    )
    data_table(
        doc,
        [
            "Block size",
            "Combined (Gb/s)",
            "Read (Gb/s)",
            "Write (Gb/s)",
            "Outstanding",
            "Notes",
        ],
        [],
    )

    doc.add_heading("LMCache bench l2 results", level=2)
    doc.add_paragraph(
        "All results below passed the corpus-verification check."
    )
    doc.add_heading(
        "1x initiator, 1x 400 GbE — per-initiator results", level=2
    )
    data_table(
        doc,
        [
            "Model",
            "Workers × in-flight",
            "Window",
            "mmgi0 (Gb/s)",
            "mmgi1 (Gb/s)",
        ],
        [
            ["Mixtral 64 KiB", "w=32 / if=8", "10 min", "67.5", "65.2"],
            ["Mixtral 256 KiB", "w=48 / if=16", "15 min", "283.1", "273.8"],
            [
                "MiniMax-M3 (run 1)",
                "w=32 / if=8 (mmgi0); w=48 / if=16 (mmgi1)",
                "15 min",
                "367.7",
                "370.2",
            ],
            [
                "MiniMax-M3 (run 2)",
                "same as run 1",
                "15 min",
                "371.6",
                "374.1",
            ],
        ],
    )
    doc.add_heading("Per-initiator operations and latency", level=2)
    data_table(
        doc,
        ["Model", "Initiator", "Throughput (Gb/s)", "Ops/s", "Avg / p99 (ms)"],
        [
            ["Mixtral 64 KiB", "mmgi0", "67.5", "128.7k", "2.44 / 3.53"],
            ["Mixtral 64 KiB", "mmgi1", "65.2", "124.3k", "2.51 / 3.75"],
            ["Mixtral 256 KiB", "mmgi0", "283.1", "135.0k", "4.52 / 6.42"],
            ["Mixtral 256 KiB", "mmgi1", "273.8", "130.5k", "4.58 / 7.01"],
            ["MiniMax-M3 (run 2)", "mmgi0", "371.6", "4.74k", "13.49 / 16.05"],
            ["MiniMax-M3 (run 2)", "mmgi1", "374.1", "4.78k", "26.80 / 35.03"],
        ],
    )
    doc.add_paragraph(
        "The MiniMax re-run reproduced the first run closely: 371.6 vs "
        "367.7 Gb/s on mmgi0 and 374.1 vs 370.2 Gb/s on mmgi1. Its "
        "mmgi1 p99 latency (35.03 ms) is higher than mmgi0's (16.05 ms), "
        "but the two initiators used different worker/in-flight settings, "
        "so this is an observation rather than evidence of a host defect."
    )
    doc.add_heading("Falcon telemetry captures", level=2)
    doc.add_paragraph(
        "Target dashboard snapshots over the measured benchmark windows."
    )
    figure(
        doc,
        DIAGRAMS_DIR / "falcon-mixtral64kb.png",
        "Mixtral 64 KiB: Falcon target telemetry during the measured run.",
    )
    figure(
        doc,
        DIAGRAMS_DIR / "falcon-mixtral256kb.png",
        "Mixtral 256 KiB: Falcon target telemetry during the measured run.",
    )
    figure(
        doc,
        DIAGRAMS_DIR / "falcon-minimax-m3.png",
        "MiniMax-M3: Falcon target telemetry during the measured run.",
    )
    doc.add_heading(
        "2x initiators, 2x 400 GbE combined — results",
        level=2,
    )
    data_table(
        doc,
        [
            "Model",
            "Combined throughput (Gb/s)",
            "DDIO absorption (avg)",
            "LLC hit % (read/write)",
            "DRAM BW (avg)",
        ],
        [
            [
                "Mixtral 64 KiB",
                "132.7",
                "95.8%",
                "99.9% / 99.0%",
                "1.5 GB/s",
            ],
            [
                "Mixtral 256 KiB",
                "556.9",
                "84.7%",
                "95.2% / 76.3%",
                "21.9 GB/s",
            ],
            [
                "MiniMax-M3 (run 1)",
                "737.9",
                "12.1%",
                "19.5% / 3.7%",
                "167.4 GB/s",
            ],
            [
                "MiniMax-M3 (run 2)",
                "745.7",
                "12.2%",
                "19.5% / 3.4%",
                "169.8 GB/s",
            ],
        ],
    )
    doc.add_paragraph(
        "The 64 KiB results illustrate why object-shaped work should not be "
        "reported as a percentage of FIO: FIO measured about 325.6 Gb/s per "
        "initiator at its own 64 KiB, 256-outstanding workload, whereas this "
        "bench l2 run reached 67.5 Gb/s on mmgi0 and 65.2 Gb/s on mmgi1 with "
        "different request structure and concurrency. Larger objects reduce "
        "per-object overhead, but MiniMax's 9.34 MiB object has no directly "
        "matching FIO row. Use FIO as a control and compare only like-sized, "
        "like-concurrency reruns taken on the same storage configuration."
    )

    doc.add_heading("5:1 mixed bench l2 results (placeholder)", level=2)
    doc.add_paragraph(
        "No 5:1 mixed-mode result is published yet. Populate this table only "
        "with a separately run `bench l2 mixed` workload."
    )
    data_table(
        doc,
        [
            "Model",
            "Combined (Gb/s)",
            "Read (Gb/s)",
            "Write (Gb/s)",
            "In-flight",
            "Notes",
        ],
        [],
    )

    doc.add_heading("Background and appendices", level=2)

    doc.add_heading("Appendix A — Setting up a host for the first time", level=3)
    doc.add_paragraph(
        "Run this only on an initiator that does not already have Python, a "
        "compiler toolchain, and the delivered LMCache checkout."
    )
    command(
        doc,
        """export http_proxy=http://proxy-dmz.intel.com:912
export https_proxy=$http_proxy
dnf install -y python3.12 python3.12-devel gcc gcc-c++ cmake make git
git clone https://github.com/tsg-/LMCache.git ~/LMCache
cd ~/LMCache
git checkout feat/bench-l2-geometry-handoff
bash scripts/ipu-poc/install_bench_l2_handoff.sh""",
    )
    doc.add_paragraph(
        "These proxy variables apply to both dnf and pip in this lab. Keep "
        "them in the same installer shell; package-manager proxy settings "
        "alone do not configure pip."
    )

    doc.add_heading("Appendix B — Setting up and tearing down storage", level=3)
    doc.add_paragraph(
        "Provision the NVMe-oF target, attach namespaces, and mount the "
        "initiator filesystem before running preflight. The benchmark does "
        "not create, format, or repair storage."
    )
    command(
        doc,
        """# On the target, before teardown: record the physical export mapping.
for n in /sys/kernel/config/nvmet/subsystems/*/namespaces/*; do
  [ -d "$n" ] || continue
  printf '%s  ' "$n"
  cat "$n/device_path"
done""",
    )
    doc.add_paragraph(
        "Record serial-backed mappings, not transient device names. Tear "
        "down top-down: stop consumers, unmount, stop the initiator array, "
        "disconnect NVMe-oF, then remove target exports. A target must not "
        "assemble an initiator-owned RAID array after a reboot."
    )

    doc.add_heading("Appendix C — Making sense of a run's results", level=3)
    doc.add_paragraph(
        "A sustained read wraps around the stored key space. Store at least "
        "two rounds with the same in-flight setting used for the timed load. "
        "A short corpus does not necessarily fail: missed keys can still "
        "inflate aggregate throughput."
    )
    doc.add_paragraph(
        "The total_success == total_keys check from Step 5 applies to every "
        "reported number here, not just the sweep's own output. The "
        "benchmark reports throughput in MiB/s despite the historical MB/s "
        "label. Increase the corpus beyond host DRAM, or use a defined "
        "cache-drop procedure, before describing the result as storage "
        "performance."
    )

    doc.add_heading("Appendix D — Optional SSD preconditioning", level=3)
    doc.add_paragraph(
        "Optional: SSD preconditioning is usually necessary for new drives "
        "only, when a reproducible, post-write-condition storage baseline is "
        "required. It writes the benchmark filesystem; confirm the mount and "
        "choose a size that fits its available capacity before running it. "
        "It is not required for functional bench l2 runs."
    )
    command(
        doc,
        """PRECONDITION_DIR=/mnt/lmcache/fio-precondition
SIZE_PER_JOB={validated-size-per-job}

mkdir -p "$PRECONDITION_DIR"
fio --name=precondition \\
    --directory="$PRECONDITION_DIR" \\
    --rw=write --bs=1m --ioengine=libaio --direct=1 --fallocate=none \\
    --numjobs=8 --iodepth=32 --size="$SIZE_PER_JOB" --group_reporting""",
    )
    doc.add_paragraph(
        "Use the mounted benchmark filesystem, never an unvalidated raw "
        "device. This creates data that may be removed only after the "
        "preconditioning run and its intended baseline measurement are "
        "complete."
    )

    doc.add_heading("Appendix E — Establishing your FIO baseline", level=3)
    doc.add_paragraph(
        "Use FIO to establish the storage and fabric envelope at the same "
        "block sizes and aggregate outstanding I/O as the peaks in "
        "\"Established bandwidth limits.\" It is a control, not a "
        "replacement for the model-shaped bench l2 sweep."
    )
    command(
        doc,
        """DIR=/mnt/lmcache/fio-base
COMMON="--directory=$DIR --filename_format=f.\\$jobnum --ioengine=libaio \
    --direct=1 --fallocate=none --numjobs=8 --size=16G --group_reporting"

fio --name=layout $COMMON --rw=write --bs=1m --iodepth=8
filefrag -v "$DIR"/f.* | grep -c unwritten || true

# block size, iodepth pairs match the peaks in the remote (fabric) table above
for cell in 64k:32 256k:64 512k:16 1m:16; do
  bs=${cell%:*}; qd=${cell#*:}
  fio --name=randread-$bs-qd$qd $COMMON --rw=randread --bs=$bs \
    --iodepth=$qd --runtime=30 --ramp_time=10 --time_based
done""",
    )
    doc.add_paragraph(
        "The layout check must print 0 (grep finds no \"unwritten\" lines, "
        "so it exits 1 -- that is the pass); any nonzero count means reads "
        "can be served as zeros without device I/O. This command uses "
        "eight jobs, so aggregate outstanding I/O is 8 x iodepth, matching "
        "the outstanding counts in the remote (fabric) bandwidth table "
        "above; the block-size/iodepth pairs do not match the local table's "
        "peaks, which use different iodepths at 64 KiB and 256 KiB. Run it "
        "on one initiator alone and, separately, on both initiators at "
        "once, since 64 KiB is the one block size where lockstep "
        "measurably underperforms solo on this rig."
    )

    doc.add_heading("Appendix F — How the telemetry pipeline works", level=3)
    doc.add_paragraph(
        "The result JSON is the record for a sweep; Prometheus and Grafana "
        "do not replace it. Set up telemetry only after a successful "
        "functional sweep, when live inspection or independent counter "
        "checks are needed. Two metrics flows, both loopback-bound and "
        "tunnelled: host counters (collector script -> .prom textfile -> "
        "node_exporter on 127.0.0.1:9100 -> SSH tunnel -> Prometheus), and "
        "bench l2's own --serve-metrics endpoint, which exists only while "
        "one benchmark process runs, so an unavailable target between "
        "cells is expected. Every node_exporter binds loopback only, so "
        "the SSH tunnel is the sole exposure and no firewall changes are "
        "needed on any test host."
    )
    figure(
        doc,
        DIAGRAMS_DIR / "brief-telemetry-pipeline.png",
        "If the dashboard looks dead, check the SSH tunnel first.",
        width_in=6.0,
    )
    doc.add_paragraph(
        "New collectors implemented for this PoC, by source file. The "
        "generic host kit runs on any test host; the two ACC/Falcon scripts "
        "are MMG-400-specific, reaching the ACC over SSH through the IMC "
        "rather than through node_exporter's own collectors."
    )
    doc.add_page_break()
    collector_table(doc)

    doc.add_heading("Appendix G — Turning on live telemetry (optional)", level=3)
    doc.add_paragraph(
        "The target's PCIe, NIC, and NUMA collectors are target-specific, so "
        "use the instrumentation README rather than the generic host "
        "installer on mmgt."
    )
    command(
        doc,
        """# ACC telemetry for the MMG-400 rig (from scripts/ipu-poc/)
IMC_PASSWORD={imc-root-password} \
    ./install_acc_stats.sh mmgi0 ':acc1:200.0.4.3'
IMC_PASSWORD={imc-root-password} \
    ./install_acc_stats.sh mmgi1 ':acc1:200.0.3.3'
IMC_PASSWORD={imc-root-password} ./install_acc_stats.sh mmgt

# Optional: persistent gRPC shadow on mmgt; dashboard stays on polling counters.
IMC_PASSWORD={imc-root-password} ACC_GRPC_SHADOW=1 \
    ./install_acc_stats.sh mmgt

# On the deployed monitoring host (mmgi0)
cd /root/lmcache-telemetry
COLLECTOR_CHECKS='' MMG_BENCH_TUNNELS=1 MMG_BENCH_INITIATORS=1 ./up.sh""",
    )
    doc.add_paragraph(
        "The ACC installer uses separate 30 s core-busy and 10 s "
        "transport-counter timers. node_exporter is loopback-only, so the "
        "SSH tunnels are the sole network exposure. The optional mmgt gRPC "
        "shadow collector keeps the IMC hop open and samples at 2 s; it is a "
        "validation path while the dashboard continues to use polling "
        "counters."
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def parse_args() -> argparse.Namespace:
    """Parse the requested document output path."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path.home()
        / "Downloads"
        / "MMG-LMCache-POC-Userguide_internal.docx",
    )
    return parser.parse_args()


def main() -> int:
    """Generate the concise operator guide."""
    args = parse_args()
    build_document(args.output)
    print(f"wrote {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
