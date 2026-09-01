# SPDX-License-Identifier: Apache-2.0
"""Build the concise bench l2 geometry-sweep delivery guide.

Needs python-docx, which is not a project dependency: `pip install python-docx`.
"""

# Standard
import argparse
from pathlib import Path

# Third Party
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def shade_cell(cell, fill: str) -> None:
    """Apply a background fill to a table cell."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def pad_cell(cell, twips: int) -> None:
    """Apply equal internal padding to a table cell."""
    margins = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        margin = OxmlElement(f"w:{side}")
        margin.set(qn("w:w"), str(twips))
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    cell._tc.get_or_add_tcPr().append(margins)


def prevent_row_split(row) -> None:
    """Keep a one-row command block on one page."""
    properties = row._tr.get_or_add_trPr()
    no_split = OxmlElement("w:cantSplit")
    properties.append(no_split)


def command(doc: Document, text: str) -> None:
    """Add a compact shaded command block."""
    table = doc.add_table(rows=1, cols=1)
    prevent_row_split(table.rows[0])
    cell = table.cell(0, 0)
    shade_cell(cell, "F1F2F4")
    pad_cell(cell, 120)

    lines = text.strip().splitlines()
    paragraph = cell.paragraphs[0]
    paragraph.style = "Command"
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run("\n".join(lines))


def build_styles(doc: Document) -> None:
    """Set the compact operator-guide type scale and rhythm."""
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, before, after in (
        ("Title", 0, 3),
        ("Subtitle", 0, 14),
        ("Heading 1", 15, 4),
        ("Heading 2", 10, 3),
    ):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    command_style = doc.styles.add_style("Command", 1)
    command_style.base_style = doc.styles["Normal"]
    command_style.font.name = "Menlo"
    command_style.font.size = Pt(8.5)
    command_style.paragraph_format.space_after = Pt(0)
    command_style.paragraph_format.line_spacing = 1.0


def build_document(output: Path) -> None:
    """Create the DOCX guide at ``output``."""
    doc = Document()
    build_styles(doc)
    for section in doc.sections:
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)

    doc.add_heading("bench l2 geometry sweep", 0)
    subtitle = doc.add_paragraph(
        "Host inventory → preflight → identity gate → sweep",
        style="Subtitle",
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph(
        "Use this guide to validate one or more NVMe-oF initiators against a "
        "target. The inventory names machines only. The coordinator rejects an "
        "unmounted or local filesystem instead of accepting a user-supplied "
        "storage path."
    )
    doc.add_heading("Execution assumptions", level=2)
    doc.add_paragraph(
        "Run the coordinator from any machine that can SSH to the inventory "
        "hosts. No controller host is configured. On each initiator, "
        "run_model_geometry.sh finds an LMCache interpreter automatically."
    )
    doc.add_paragraph(
        "Before you run preflight, confirm that the coordinator can use "
        "passwordless SSH to both initiators and the target, that each "
        "initiator has the delivered LMCache checkout, and that the target "
        "namespaces are attached and mounted on the initiators."
    )

    doc.add_heading("1. Use the supplied host inventory", level=1)
    doc.add_paragraph(
        "Use the supplied inventory file "
        "scripts/ipu-poc/inventories/mmg-two-initiator.env as the argument to "
        "every coordinator command. It names machines only. Edit its host "
        "aliases only when running on a different rig; do not put IP "
        "addresses, filesystem paths, Python paths, or workload settings in "
        "this file."
    )
    command(
        doc,
        """# scripts/ipu-poc/inventories/mmg-two-initiator.env
INITIATOR_HOSTS=(mmgi0 mmgi1)
TARGET_HOST=mmgt""",
    )
    doc.add_paragraph(
        "The coordinator runs the benchmark on each initiator. The target "
        "remains storage-only; it is contacted only to verify its live "
        "NVMe-oF export addresses."
    )
    doc.add_heading("2. Run preflight", level=1)
    command(
        doc,
        """bash scripts/ipu-poc/run_geometry_inventory.sh preflight \\
  scripts/ipu-poc/inventories/mmg-two-initiator.env""",
    )
    doc.add_paragraph(
        "Preflight checks the checkout, interpreter, mount, and NVMe-oF "
        "controller on every initiator. It stops before writing anything."
    )

    doc.add_heading("3. Verify corpus identity", level=1)
    command(
        doc,
        """RUN_ID=mmg-verify-$(date +%Y%m%d-%H%M%S) \\
  bash scripts/ipu-poc/run_geometry_inventory.sh verify \\
  scripts/ipu-poc/inventories/mmg-two-initiator.env""",
    )
    doc.add_paragraph(
        "The verification gate writes a small corpus, requires the storing "
        "profile to read it back, and requires a same-geometry mismatch to "
        "read zero objects. It is a correctness check, not a bandwidth result."
    )

    doc.add_page_break()
    doc.add_heading("4. Run the sweep", level=1)
    command(
        doc,
        """export RUN_ID=mmg-geometry-$(date +%Y%m%d-%H%M%S)
IN_FLIGHTS='8 16 24' WARMUP_SEC=60 DURATION_SEC=120 INCLUDE_MINIMAX=1 \\
  bash scripts/ipu-poc/run_geometry_inventory.sh sweep \\
  scripts/ipu-poc/inventories/mmg-two-initiator.env""",
    )
    doc.add_paragraph(
        "For every initiator, the coordinator runs the five page-burst "
        "profiles, in page-size order: Mixtral 64 KiB, Mixtral 128 KiB, "
        "DeepSeek-V3 144 KiB, Mixtral 256 KiB, and Mixtral 512 KiB, then "
        "MiniMax-M3. Every profile is stored and read at 8, 16, and 24 "
        "in-flight submits. Each sustained load warms up for 60 seconds and "
        "measures for 120 seconds. MiniMax-M3 uses the O_DIRECT-compatible "
        "padding profile, which adds a 3,072 B alignment tail (about 0.03% "
        "of an object). Its fp8 DSA-indexer geometry assumption remains to "
        "be confirmed, so retain that caveat when comparing it with the "
        "five page-burst profiles. Each host writes only below "
        "BENCH_MOUNT/bench-l2/<run-id>/<hostname>, which defaults to "
        "/mnt/lmcache, and uses a host-specific key prefix. Result JSON "
        "files remain in the remote checkout's results/ directory and "
        "include the in-flight count in their names."
    )

    doc.add_heading("5. Accept or reject the result", level=1)
    doc.add_paragraph(
        "For each result JSON, require total_success == total_keys. A miss "
        "invalidates the bandwidth figure. The default corpus is small enough "
        "to fit in DRAM on most benchmark hosts, so it is a functional sweep; "
        "scale the corpus beyond memory before making a storage-performance "
        "claim."
    )
    doc.add_heading("Collect the result JSON", level=2)
    doc.add_paragraph(
        "Keep the exported RUN_ID from step 4. Use the supplied inventory to "
        "list every result produced by this sweep; copy the JSON files off the "
        "initiators before comparing or sharing a result."
    )
    command(
        doc,
        """source scripts/ipu-poc/inventories/mmg-two-initiator.env
for host in "${INITIATOR_HOSTS[@]}"; do
  ssh "$host" bash -s -- "$RUN_ID" <<'REMOTE'
set -euo pipefail
run_id=$1
for repo in "$HOME/LMCache" /root/LMCache; do
  [ -x "$repo/scripts/ipu-poc/run_model_geometry.sh" ] && break
done
[ -x "$repo/scripts/ipu-poc/run_model_geometry.sh" ] ||
  { echo "ABORT: LMCache geometry checkout not found" >&2; exit 2; }
find "$repo/results" -maxdepth 1 -name "$run_id-*.json" -print
REMOTE
done""",
    )

    doc.add_heading("Background and appendices", level=1)

    doc.add_heading("Why corpus verification matters", level=2)
    doc.add_paragraph(
        "The profile is part of corpus identity. A store under one profile "
        "must not be accepted by another profile that happens to use the same "
        "page size. The verification command stores a Mixtral corpus, loads it with "
        "the storing profile, then requires zero hits from a byte-different "
        "profile with the same geometry."
    )
    doc.add_paragraph(
        "That is a correctness gate, not a performance result. A passed gate "
        "means the benchmark can distinguish the intended corpus on the "
        "filesystem under test."
    )

    doc.add_heading("Appendix A — First-time host setup", level=2)
    doc.add_paragraph(
        "Run this only on an initiator that does not already have Python, a "
        "compiler toolchain, and the delivered LMCache checkout."
    )
    command(
        doc,
        """dnf install -y python3.12 python3.12-devel gcc gcc-c++ cmake make git
cd LMCache
bash scripts/ipu-poc/install_bench_l2_handoff.sh""",
    )
    doc.add_paragraph(
        "If the host requires a proxy, export http_proxy and https_proxy in "
        "the installer shell. Package-manager proxy settings do not configure "
        "pip."
    )

    doc.add_heading("Appendix B — Storage lifecycle", level=2)
    doc.add_paragraph(
        "Provision the NVMe-oF target, attach namespaces, and mount the "
        "initiator filesystem before running preflight. The benchmark does "
        "not create, format, or repair storage."
    )
    command(
        doc,
        """# On the target, before teardown: record the physical export mapping.
for n in /sys/kernel/config/nvmet/subsystems/*/namespaces/*; do
  [ -d "$n" ] || continue
  printf '%s  ' "$n"
  cat "$n/device_path"
done""",
    )
    doc.add_paragraph(
        "Record serial-backed mappings, not transient device names. Teardown "
        "runs from the top down: stop consumers, unmount, stop the initiator "
        "array, disconnect NVMe-oF, then remove target exports. A target must "
        "not assemble an initiator-owned RAID array after a reboot."
    )

    doc.add_heading("Appendix C — Interpreting a run", level=2)
    doc.add_paragraph(
        "A sustained read wraps around the stored key space. Store at least "
        "two rounds with the same in-flight setting used for the timed load. "
        "A short corpus does not necessarily fail: missed keys can still "
        "inflate aggregate throughput."
    )
    doc.add_paragraph(
        "Treat throughput as valid only when total_success equals total_keys. "
        "The benchmark reports throughput in MiB/s despite the historical "
        "MB/s label. Increase the corpus beyond host DRAM, or use a defined "
        "cache-drop procedure, before describing the result as storage "
        "performance."
    )

    doc.add_heading("Appendix D — FIO control", level=2)
    doc.add_paragraph(
        "Use FIO to establish the storage and fabric envelope at comparable "
        "block size and aggregate outstanding I/O. It is a control, not a "
        "replacement for the model-shaped bench l2 sweep."
    )
    command(
        doc,
        """DIR=/mnt/lmcache-stage2/fio-base
COMMON="--directory=$DIR --filename_format=f.\\$jobnum --ioengine=libaio \
  --direct=1 --fallocate=none --numjobs=8 --size=16G --group_reporting"

fio --name=layout $COMMON --rw=write --bs=1m --iodepth=8
filefrag -v "$DIR/f.0" | grep -c unwritten

for qd in 4 16 64 128; do
  fio --name=randread-qd$qd $COMMON --rw=randread --bs=64k \
    --iodepth=$qd --runtime=30 --ramp_time=10 --time_based
done""",
    )
    doc.add_paragraph(
        "The layout check must report zero unwritten extents; otherwise reads "
        "can be served as zeros without device I/O. Compare a bench result "
        "with the FIO cell nearest its actual concurrency, not the FIO peak. "
        "The 64 KiB FIO cells control only the 64 KiB profile. This command "
        "uses eight jobs, so aggregate outstanding I/O is 8 × iodepth; use a "
        "matching block size and aggregate concurrency before comparing any "
        "other profile."
    )

    doc.add_heading("Appendix E — Optional telemetry setup", level=2)
    doc.add_paragraph(
        "The result JSON is the record for a sweep. Set up telemetry only "
        "after a successful functional sweep when live inspection or "
        "independent counter checks are needed; it is not a benchmark "
        "prerequisite. Prometheus and Grafana do not replace the completed-run "
        "JSON."
    )
    command(
        doc,
        """# ACC telemetry for the MMG-400 rig (from scripts/ipu-poc/)
IMC_PASSWORD=<imc-root-password> \
  ./install_acc_stats.sh mmgi0 ':acc1:200.0.4.3'
IMC_PASSWORD=<imc-root-password> \
  ./install_acc_stats.sh mmgi1 ':acc1:200.0.3.3'
IMC_PASSWORD=<imc-root-password> ./install_acc_stats.sh mmgt

# Optional: persistent gRPC shadow on mmgt; dashboard stays on polling counters.
IMC_PASSWORD=<imc-root-password> ACC_GRPC_SHADOW=1 \
  ./install_acc_stats.sh mmgt

# On the control machine
cd docs/design/v1/platform/ipu-poc/instrumentation
HOSTS="mmgt:19106 mmgi0:19107 mmgi1:19108" \
  MMG_BENCH_TUNNELS=1 MMG_BENCH_INITIATORS=1 ./up.sh""",
    )
    doc.add_paragraph(
        "The target's PCIe, NIC, and NUMA collectors are target-specific; use "
        "the instrumentation README rather than the generic host installer on "
        "mmgt. A bench metrics endpoint exists only while its CLI process runs, "
        "so an unavailable target between cells is expected. The ACC installer "
        "uses separate 30 s core-busy and 10 s transport-counter timers. The "
        "optional mmgt gRPC shadow collector keeps the IMC hop open and samples "
        "at 5 s; it is validated against the polling counters before becoming "
        "a dashboard source."
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def parse_args() -> argparse.Namespace:
    """Parse the requested document output path."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path.home() / "Downloads" / "bench-l2-geometry-code-delivery.docx",
    )
    return parser.parse_args()


def main() -> int:
    """Generate the concise operator guide."""
    args = parse_args()
    build_document(args.output)
    print(f"wrote {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
